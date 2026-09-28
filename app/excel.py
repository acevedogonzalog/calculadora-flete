"""Exportación de viajes a Excel.

Los resultados se escriben como valores (no fórmulas) porque son el registro
de cómo quedó el viaje al guardarlo, y así se ven bien en cualquier visor,
incluido el celular.
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = "Arial"
MONEY = '"$" #,##0.00;[Red]-"$" #,##0.00;"-"'
NUM = '#,##0.00;-#,##0.00;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
VERDE = "1E6B4F"
thin = Side(style="thin", color="C9CFCB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F_TIT = Font(name=F, size=15, bold=True, color="FFFFFF")
F_SEC = Font(name=F, size=11, bold=True, color="FFFFFF")
F_LBL = Font(name=F, size=11)
F_VAL = Font(name=F, size=11)
F_BOLD = Font(name=F, size=11, bold=True)
FILL_TIT = PatternFill("solid", fgColor="1B2220")
FILL_SEC = PatternFill("solid", fgColor=VERDE)
FILL_RES = PatternFill("solid", fgColor="E3EFE9")


def _fecha(v):
    return date.fromisoformat(v) if v else None


def _div(a, b):
    return a / b if b else None


def excel_viaje(v: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Viaje"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 30
    fila = 2

    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=3)
    c = ws.cell(fila, 2, f"VIAJE Nº {v['id']} · {v.get('cliente') or 'Sin cliente'}")
    c.font, c.fill = F_TIT, FILL_TIT
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[fila].height = 28
    fila += 2

    def seccion(titulo):
        nonlocal fila
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=3)
        c = ws.cell(fila, 2, titulo)
        c.font, c.fill = F_SEC, FILL_SEC
        c.alignment = Alignment(indent=1)
        fila += 1

    def linea(etq, valor, fmt=None, destacar=False):
        nonlocal fila
        a = ws.cell(fila, 2, etq)
        b = ws.cell(fila, 3, valor)
        a.font = F_BOLD if destacar else F_LBL
        b.font = F_BOLD if destacar else F_VAL
        a.border = b.border = BOX
        a.alignment = Alignment(indent=1)
        if fmt:
            b.number_format = fmt
        if isinstance(valor, (int, float)) or fmt:
            b.alignment = Alignment(horizontal="right")
        if destacar:
            a.fill = b.fill = FILL_RES
        fila += 1

    seccion("FACTURA")
    linea("Fecha", _fecha(v.get("fecha")), "dd/mm/yyyy")
    linea("Nº de factura", v.get("factura_nro") or "")
    linea("Cliente", v.get("cliente") or "")
    linea("CUIT cliente", v.get("cliente_cuit") or "")
    linea("Detalle", v.get("detalle") or "")
    linea("Remito", v.get("remito") or "")
    linea("CAE", v.get("cae") or "")
    fila += 1

    seccion("DATOS DEL VIAJE")
    linea("Toneladas", v["toneladas"], NUM)
    linea("Precio por tonelada", v["precio_tn"], MONEY)
    linea("Porcentaje chofer", v["pct_chofer"] / 100, PCT)
    linea("Precio gasoil por litro", v["precio_gasoil"], MONEY)
    linea("Litros de gasoil", v["litros"], NUM)
    linea("Kilómetros", v.get("km"), NUM)
    linea("Otros gastos", v["otros_gastos"], MONEY)
    fila += 1

    seccion("RESULTADOS")
    linea("Total factura", v["total"], MONEY, destacar=True)
    linea("Pago al chofer", v["pago_chofer"], MONEY)
    linea("Costo de combustible", v["costo_gasoil"], MONEY)
    linea("Otros gastos", v["otros_gastos"], MONEY)
    linea("Total gastos", v["total_gastos"], MONEY, destacar=True)
    linea("Ganancia neta", v["ganancia"], MONEY, destacar=True)
    fila += 1

    seccion("INDICADORES")
    km = v.get("km") or 0
    linea("Margen de ganancia", _div(v["ganancia"], v["total"]), PCT)
    linea("Combustible sobre la factura", _div(v["costo_gasoil"], v["total"]), PCT)
    linea("Ganancia por tonelada", _div(v["ganancia"], v["toneladas"]), MONEY)
    linea("Gasoil por tonelada", _div(v["costo_gasoil"], v["toneladas"]), MONEY)
    linea("Consumo (L cada 100 km)", _div(v["litros"] * 100, km), NUM)
    linea("Costo por km", _div(v["total_gastos"], km), MONEY)
    linea("Ganancia por km", _div(v["ganancia"], km), MONEY)

    if v.get("notas"):
        fila += 1
        seccion("NOTAS")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=3)
        c = ws.cell(fila, 2, v["notas"])
        c.font = F_VAL
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[fila].height = max(30, 15 * (1 + len(v["notas"]) // 60))

    return _guardar(wb)


COLUMNAS = [
    ("Nº", "id", None, 6),
    ("Fecha", "fecha", "dd/mm/yyyy", 12),
    ("Factura", "factura_nro", None, 17),
    ("Cliente", "cliente", None, 26),
    ("Detalle", "detalle", None, 28),
    ("Remito", "remito", None, 16),
    ("Toneladas", "toneladas", NUM, 11),
    ("Precio x tn", "precio_tn", MONEY, 14),
    ("Total factura", "total", MONEY, 16),
    ("% Chofer", "pct_chofer", PCT, 10),
    ("Pago chofer", "pago_chofer", MONEY, 15),
    ("Gasoil $/L", "precio_gasoil", MONEY, 12),
    ("Litros", "litros", NUM, 10),
    ("Costo gasoil", "costo_gasoil", MONEY, 15),
    ("Km", "km", NUM, 10),
    ("Otros gastos", "otros_gastos", MONEY, 14),
    ("Total gastos", "total_gastos", MONEY, 15),
    ("Ganancia neta", "ganancia", MONEY, 16),
    ("Notas", "notas", None, 30),
]
SUMAR = {"toneladas", "total", "pago_chofer", "litros", "costo_gasoil", "km",
         "otros_gastos", "total_gastos", "ganancia"}


def excel_todos(lista: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Viajes"
    ws.freeze_panes = "A2"
    for col, (titulo, _, _, ancho) in enumerate(COLUMNAS, 1):
        c = ws.cell(1, col, titulo)
        c.font, c.fill, c.border = F_SEC, FILL_SEC, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[1].height = 30

    for r, v in enumerate(lista, 2):
        for col, (_, clave, fmt, _) in enumerate(COLUMNAS, 1):
            val = v.get(clave)
            if clave == "fecha":
                val = _fecha(val)
            elif clave == "pct_chofer":
                val = val / 100
            c = ws.cell(r, col, val)
            c.font, c.border = F_VAL, BOX
            if fmt:
                c.number_format = fmt

    tr = len(lista) + 2
    ws.cell(tr, 1, "TOTAL")
    for col, (_, clave, fmt, _) in enumerate(COLUMNAS, 1):
        c = ws.cell(tr, col)
        c.font, c.fill, c.border = F_BOLD, FILL_RES, BOX
        if clave in SUMAR:
            c.value = sum((v.get(clave) or 0) for v in lista)
            c.number_format = fmt
    return _guardar(wb)


def _guardar(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
