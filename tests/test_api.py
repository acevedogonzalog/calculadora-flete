"""Tests de punta a punta de la API: uso normal, valores extremos y ataques."""
import io
import json

import pytest
from openpyxl import load_workbook

from conftest import VIAJE_OK, jpeg


def crear(client, **cambios):
    r = client.post("/api/viajes", json={**VIAJE_OK, **cambios})
    assert r.status_code == 201, r.text
    return r.json()


# ---------------- Uso normal ----------------

def test_estado_sin_clave(client):
    assert client.get("/api/estado").json() == {"requiere_clave": False, "autenticado": True}


def test_pagina_principal(client):
    r = client.get("/")
    assert r.status_code == 200 and "Calculadora de Flete" in r.text


def test_factura_real_se_lee_y_se_guarda(client, factura_pdf):
    r = client.post("/api/factura", files={"archivo": ("f.pdf", factura_pdf, "application/pdf")})
    d = r.json()
    assert r.status_code == 200
    assert d["toneladas"] == 28 and d["precio_tn"] == 54049 and d["importe_total"] == 1513372
    assert d["cliente"] == "SERVAGROP HDO S. A." and d["fecha"] == "2026-09-28"
    assert d["factura_nro"] == "00001-00000106" and d["cae"] == "86395043581706"
    assert d["factura_archivo"].endswith(".pdf") and d["avisos"] == []


def test_calculo_del_viaje(client):
    v = crear(client)
    assert v["total"] == 1513372
    assert v["pago_chofer"] == pytest.approx(272406.96)
    assert v["costo_gasoil"] == 575000
    assert v["ganancia"] == pytest.approx(665965.04)


def test_viaje_con_factura_y_remitos_completo(client, factura_pdf):
    f = client.post("/api/factura", files={"archivo": ("f.pdf", factura_pdf, "application/pdf")}).json()
    r1 = client.post("/api/remitos", files={"archivo": ("a.jpg", jpeg(4000, 3000, 6), "image/jpeg")}).json()
    r2 = client.post("/api/remitos", files={"archivo": ("b.pdf", factura_pdf, "application/pdf")}).json()
    v = crear(client, factura_archivo=f["factura_archivo"], factura_nombre="f.pdf", remitos=[r1, r2])
    assert v["remitos_cant"] == 2
    pdf = client.get(f"/api/viajes/{v['id']}/factura")
    assert pdf.status_code == 200 and pdf.content == factura_pdf
    foto = client.get(f"/api/remitos/archivo/{r1['archivo']}")
    from PIL import Image
    im = Image.open(io.BytesIO(foto.content))
    assert im.size == (1800, 2400)  # enderezada (EXIF 6) y reducida
    assert client.get(f"/api/remitos/archivo/{r1['miniatura']}").status_code == 200


def test_exportaciones_excel(client):
    v = crear(client)
    for url in (f"/api/viajes/{v['id']}/excel", "/api/viajes/excel", "/api/resumen/excel?mes=2026-09"):
        r = client.get(url)
        assert r.status_code == 200, url
        load_workbook(io.BytesIO(r.content))  # abre sin errores


def test_resumen_mensual(client):
    crear(client, fecha="2026-07-31")
    crear(client, fecha="2026-08-01")
    r = client.get("/api/resumen?mes=2026-08").json()
    assert all(v["fecha"].startswith("2026-08") for v in r["viajes"])
    assert r["anterior"]["mes"] == "2026-07"
    assert client.get("/api/resumen?mes=2026-13").status_code == 400
    assert client.get("/api/resumen?mes=1999-01").json()["viajes"] == []


def test_resumen_mes_vacio_no_divide_por_cero(client):
    r = client.get("/api/resumen?mes=2000-01")
    assert r.status_code == 200 and r.json()["totales"]["margen"] is None


def test_borrar_viaje_borra_sus_remitos(client):
    foto = client.post("/api/remitos", files={"archivo": ("a.jpg", jpeg(), "image/jpeg")}).json()
    v = crear(client, remitos=[foto])
    assert client.delete(f"/api/viajes/{v['id']}").status_code == 200
    assert client.get(f"/api/viajes/{v['id']}").status_code == 404
    assert client.delete(f"/api/viajes/{v['id']}").status_code == 404


def test_adjuntar_y_quitar_remito_despues(client):
    v = crear(client)
    d = client.post(f"/api/viajes/{v['id']}/remitos", files={"archivo": ("a.jpg", jpeg(), "image/jpeg")}).json()
    assert len(d["remitos"]) == 1
    assert client.delete(f"/api/remitos/{d['remitos'][0]['id']}").status_code == 200
    assert client.get(f"/api/viajes/{v['id']}").json()["remitos"] == []


# ---------------- Valores inválidos o extremos ----------------

@pytest.mark.parametrize("cambios", [
    {"toneladas": -1}, {"pct_chofer": 101}, {"litros": -5}, {"km": -1},
    {"cliente": "x" * 201}, {"notas": "x" * 2001}, {"fecha": "2026-02-30"},
    {"toneladas": "abc"}, {"remitos": [{"archivo": "a.jpg"}] * 31},
])
def test_rechaza_datos_invalidos(client, cambios):
    assert client.post("/api/viajes", json={**VIAJE_OK, **cambios}).status_code == 422


def test_toneladas_cero_se_rechaza(client):
    assert client.post("/api/viajes", json={**VIAJE_OK, "toneladas": 0}).status_code == 400


@pytest.mark.parametrize("valor", ["Infinity", "NaN", "1e308"])
def test_numeros_infinitos_o_gigantes_no_rompen_la_lista(client, valor):
    cuerpo = json.dumps({**VIAJE_OK, "precio_tn": 1}).replace('"precio_tn": 1', f'"precio_tn": {valor}')
    r = client.post("/api/viajes", content=cuerpo, headers={"content-type": "application/json"})
    assert r.status_code == 422, f"{valor} fue aceptado"
    assert client.get("/api/viajes").status_code == 200
    assert client.get("/api/resumen").status_code == 200


@pytest.mark.parametrize("vid", ["2147483648", "99999999999999999999", "-1", "abc"])
def test_ids_fuera_de_rango(client, vid):
    assert client.get(f"/api/viajes/{vid}").status_code in (404, 422)


def test_ruta_api_inexistente_es_404(client):
    assert client.get("/api/no-existe").status_code == 404


# ---------------- Archivos maliciosos o inválidos ----------------

def test_factura_que_no_es_pdf(client):
    r = client.post("/api/factura", files={"archivo": ("x.pdf", b"<html><script>alert(1)</script>", "application/pdf")})
    assert r.status_code == 400


def test_factura_pdf_basura_no_rompe(client):
    r = client.post("/api/factura", files={"archivo": ("x.pdf", b"%PDF-1.4 basura" * 10, "application/pdf")})
    assert r.status_code == 200 and r.json()["avisos"]


def test_factura_demasiado_grande(client):
    r = client.post("/api/factura", files={"archivo": ("x.pdf", b"%PDF" + b"0" * (11 * 1024 * 1024), "application/pdf")})
    assert r.status_code == 413


@pytest.mark.parametrize("contenido", [b"", b"hola", b"GIF89a-no-es-imagen", b"<svg onload=alert(1)></svg>"])
def test_remito_que_no_es_imagen(client, contenido):
    r = client.post("/api/remitos", files={"archivo": ("x.jpg", contenido, "image/jpeg")})
    assert r.status_code == 400


def test_bomba_de_descompresion(client):
    """PNG chiquito que al abrirlo ocupa gigas de memoria: tiene que rechazarse, no colgar el servidor."""
    from PIL import Image
    buf = io.BytesIO()
    Image.new("1", (30000, 30000)).save(buf, "PNG")
    assert len(buf.getvalue()) < 2_000_000
    r = client.post("/api/remitos", files={"archivo": ("bomba.png", buf.getvalue(), "image/png")})
    assert r.status_code == 400


@pytest.mark.parametrize("nombre", [
    "..%2F..%2Fetc%2Fpasswd", "..%2Ftest.db", "%2Fetc%2Fpasswd", "x.jpg%00.pdf",
    "..%5C..%5Cwindows", "remito.jpg%2F..%2F..%2Ftest.db",
])
def test_path_traversal_en_remitos(client, nombre):
    r = client.get(f"/api/remitos/archivo/{nombre}")
    assert r.status_code == 404
    assert b"root:" not in r.content and b"SQLite" not in r.content


def test_path_traversal_en_pagina(client):
    for ruta in ("/../app/server.py", "/..%2Fapp%2Fserver.py", "/%2e%2e/app/server.py", "/../../etc/passwd"):
        r = client.get(ruta)
        assert b"CLAVE" not in r.content and b"root:" not in r.content, ruta


def test_referencias_falsas_a_archivos_se_ignoran(client):
    v = crear(client, factura_archivo="../../etc/passwd.pdf",
              remitos=[{"archivo": "../test.db"}, {"archivo": "noexiste.jpg"}])
    assert v["factura_archivo"] == "" and v["remitos"] == []


def test_nombre_de_archivo_malicioso_en_descarga(client, factura_pdf):
    """El nombre original del archivo no puede romper la cabecera Content-Disposition."""
    r = client.post("/api/factura", files={"archivo": ('a"; x=1\r\nSet-Cookie: pwn=1.pdf', factura_pdf, "application/pdf")})
    d = r.json()
    v = crear(client, factura_archivo=d["factura_archivo"], factura_nombre=d["factura_nombre"])
    resp = client.get(f"/api/viajes/{v['id']}/factura")
    assert "pwn" not in resp.cookies and '"; x=' not in resp.headers["content-disposition"]


# ---------------- Excel: inyección de fórmulas ----------------

@pytest.mark.parametrize("texto", ['=HYPERLINK("http://malo.com","Clic")', "+1+1", "-2+3", "@SUM(1)", "=cmd|' /C calc'!A0"])
def test_excel_no_ejecuta_formulas_de_los_datos(client, texto):
    v = crear(client, cliente=texto, notas=texto, detalle=texto)
    for url in (f"/api/viajes/{v['id']}/excel", "/api/viajes/excel"):
        wb = load_workbook(io.BytesIO(client.get(url).content))
        for ws in wb.worksheets:
            for fila in ws.iter_rows():
                for c in fila:
                    if isinstance(c.value, str) and texto.lstrip("'") in c.value:
                        assert c.data_type != "f", f"{url}: '{texto}' quedó como fórmula en {c.coordinate}"


# ---------------- Cabeceras de seguridad ----------------

def test_cabeceras_de_seguridad(client):
    for url in ("/", "/api/viajes", "/api/no-existe"):
        h = client.get(url).headers
        assert h.get("x-content-type-options") == "nosniff", url
        assert h.get("x-frame-options") == "DENY", url
    assert "frame-ancestors 'none'" in client.get("/").headers.get("content-security-policy", "")
    assert client.get("/api/viajes").headers.get("cache-control") == "no-store"


def test_error_de_validacion_es_claro(client):
    r = client.post("/api/viajes", json={**VIAJE_OK, "litros": -1})
    assert r.status_code == 422 and "Litros" in r.json()["detail"]


def test_subida_con_content_length_enorme(client):
    r = client.post("/api/remitos", content=b"x", headers={
        "content-type": "multipart/form-data; boundary=abc", "content-length": str(500 * 1024 * 1024)})
    assert r.status_code == 413
